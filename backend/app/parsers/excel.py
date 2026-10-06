from __future__ import annotations

import io
from app.parsers.base import DocumentBlock


class ExcelParser:
    async def parse(
        self,
        file_bytes: bytes,
        filename: str,
        mime_type: str | None = None,
        storage_path: str | None = None,
    ) -> list[DocumentBlock]:
        try:
            import openpyxl
        except ImportError:
            # Fallback to plain-text decode if openpyxl is not present
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="ignore")
            return [DocumentBlock(type="text", content=text, page_number=1)]

        blocks = []
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]

                # Extract all rows from the sheet
                rows = []
                for r in sheet.iter_rows(values_only=True):
                    # Filter out completely empty rows
                    if any(c is not None for c in r):
                        rows.append([str(c).strip() if c is not None else "" for c in r])

                if not rows:
                    continue

                headers = rows[0]
                if not headers or all(h == "" for h in headers):
                    continue

                data_rows = rows[1:]
                if not data_rows:
                    markdown_table = f"| {' | '.join(headers)} |\n| {' | '.join(['---'] * len(headers))} |"
                    blocks.append(
                        DocumentBlock(
                            type="table",
                            content=markdown_table,
                            page_number=1,
                            media_path=storage_path,
                            metadata={"sheet_name": sheet_name},
                        )
                    )
                    continue

                chunk_size = 10
                for idx in range(0, len(data_rows), chunk_size):
                    chunk_data = data_rows[idx : idx + chunk_size]

                    lines = []
                    # Include sheet name as context for search queries
                    lines.append(f"Sheet: {sheet_name}")
                    lines.append("| " + " | ".join(headers) + " |")
                    lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                    for r in chunk_data:
                        padded_row = list(r)
                        if len(padded_row) < len(headers):
                            padded_row += [""] * (len(headers) - len(padded_row))
                        elif len(padded_row) > len(headers):
                            padded_row = padded_row[: len(headers)]
                        lines.append("| " + " | ".join(padded_row) + " |")

                    markdown_table = "\n".join(lines)
                    blocks.append(
                        DocumentBlock(
                            type="table",
                            content=markdown_table,
                            page_number=1,
                            media_path=storage_path,
                            metadata={"sheet_name": sheet_name},
                        )
                    )
            wb.close()
        except Exception:
            pass

        return blocks
