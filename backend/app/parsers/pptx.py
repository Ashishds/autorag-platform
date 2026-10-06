from __future__ import annotations

import io

from app.parsers.base import DocumentBlock


class PPTXParser:
    async def parse(
        self,
        file_bytes: bytes,
        filename: str,
        mime_type: str | None = None,
        storage_path: str | None = None,
    ) -> list[DocumentBlock]:
        blocks: list[DocumentBlock] = []
        try:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(file_bytes))
            for slide_num, slide in enumerate(prs.slides, start=1):
                texts: list[str] = []
                for shape in slide.shapes:
                    text = getattr(shape, "text", "") or ""
                    if text.strip():
                        texts.append(text.strip())
                if texts:
                    blocks.append(
                        DocumentBlock(
                            type="text",
                            content="\n".join(texts),
                            page_number=slide_num,
                            metadata={"slide": slide_num},
                        )
                    )
        except ImportError:
            raise ValueError(
                "PowerPoint support requires python-pptx. Install with: uv add python-pptx"
            ) from None
        except Exception as exc:
            raise ValueError(f"Failed to parse PowerPoint file: {exc}") from exc

        return blocks
